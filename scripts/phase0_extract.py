from __future__ import annotations

import argparse
import hashlib
import json
import re
import zipfile
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path, PurePosixPath

import pdfplumber
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


XML_NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def normalize_text(value: str) -> str:
    value = value.replace("\u200c", " ").replace("\ufeff", " ")
    value = re.sub(r"\s+", " ", value).strip()
    return value


def safe_text(value: object) -> str:
    return normalize_text(str(value or ""))


def rgb_to_hex(color: object) -> str | None:
    try:
        rgb = color.rgb
        return str(rgb) if rgb is not None else None
    except (AttributeError, TypeError, ValueError):
        return None


def paragraph_payload(paragraph) -> dict:
    runs = []
    for run in paragraph.runs:
        font = run.font
        runs.append(
            {
                "text": run.text,
                "bold": font.bold,
                "italic": font.italic,
                "sizePt": round(font.size.pt, 2) if font.size else None,
                "fontName": font.name,
                "color": rgb_to_hex(font.color),
            }
        )
    return {
        "text": normalize_text(paragraph.text),
        "level": paragraph.level,
        "alignment": str(paragraph.alignment) if paragraph.alignment else None,
        "runs": runs,
    }


def text_frame_payload(text_frame) -> dict:
    paragraphs = [paragraph_payload(p) for p in text_frame.paragraphs]
    return {
        "text": normalize_text("\n".join(p["text"] for p in paragraphs if p["text"])),
        "paragraphs": paragraphs,
        "verticalAnchor": str(text_frame.vertical_anchor) if text_frame.vertical_anchor else None,
    }


def shape_payload(shape, slide_width: int, slide_height: int, parent_id: str | None = None) -> list[dict]:
    payload = {
        "shapeId": str(shape.shape_id),
        "parentShapeId": parent_id,
        "name": shape.name,
        "shapeType": str(shape.shape_type),
        "bbox": {
            "x": round(shape.left / slide_width, 6),
            "y": round(shape.top / slide_height, 6),
            "w": round(shape.width / slide_width, 6),
            "h": round(shape.height / slide_height, 6),
        },
        "rotation": round(shape.rotation or 0, 2),
        "text": None,
        "table": None,
        "image": None,
        "isConnector": shape.shape_type == MSO_SHAPE_TYPE.LINE,
    }

    if getattr(shape, "has_text_frame", False):
        payload["text"] = text_frame_payload(shape.text_frame)

    if getattr(shape, "has_table", False):
        rows = []
        for row in shape.table.rows:
            rows.append([normalize_text(cell.text) for cell in row.cells])
        payload["table"] = {"rows": rows, "rowCount": len(rows), "columnCount": len(rows[0]) if rows else 0}

    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        try:
            image = shape.image
            payload["image"] = {
                "filename": image.filename,
                "contentType": image.content_type,
                "sha1": image.sha1,
                "size": list(image.size),
                "linked": False,
            }
        except ValueError:
            blip = shape._element.xpath(".//a:blip")
            embedded_id = blip[0].get(f"{{{XML_NS['r']}}}embed") if blip else None
            linked_id = blip[0].get(f"{{{XML_NS['r']}}}link") if blip else None
            payload["image"] = {
                "filename": None,
                "contentType": None,
                "sha1": None,
                "size": None,
                "linked": bool(linked_id),
                "embeddedRelationshipId": embedded_id,
                "linkedRelationshipId": linked_id,
                "unsupportedByPythonPptx": True,
            }

    shapes = [payload]
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for child in shape.shapes:
            shapes.extend(shape_payload(child, slide_width, slide_height, str(shape.shape_id)))
    return shapes


def largest_text_title(shapes: list[dict]) -> str:
    candidates = []
    for shape in shapes:
        text = (shape.get("text") or {}).get("text") or ""
        if not text:
            continue
        sizes = [r.get("sizePt") for p in (shape.get("text") or {}).get("paragraphs", []) for r in p.get("runs", []) if r.get("sizePt")]
        size = max(sizes) if sizes else 0
        bbox = shape["bbox"]
        candidates.append((size, -bbox["y"], -len(text), text))
    if not candidates:
        return ""
    return max(candidates)[-1].split("\n", 1)[0][:180]


def infer_title(shapes: list[dict]) -> str:
    candidates = []
    ignored = {"---", "ilia", "iliA", "مقدمه"}
    for shape in shapes:
        text = ((shape.get("text") or {}).get("text") or "").strip()
        first_line = text.split("\n", 1)[0].strip()
        if not first_line or first_line.casefold() in {value.casefold() for value in ignored}:
            continue
        if len(first_line) < 3 or re.fullmatch(r"[\d۰-۹٠-٩./()\-]+", first_line):
            continue
        bbox = shape["bbox"]
        if bbox["y"] <= 0.13:
            candidates.append((bbox["y"], -bbox["x"], -len(first_line), first_line))
    if candidates:
        return min(candidates)[-1][:180]
    return largest_text_title(shapes)


def classify_slide(slide: dict) -> str:
    text = slide["text"]
    shape_count = slide["metrics"]["shapeCount"]
    image_count = slide["metrics"]["imageCount"]
    table_count = slide["metrics"]["tableCount"]
    if not text and image_count:
        return "visual-only"
    if not text:
        return "empty-or-background-only"
    if table_count:
        return "table-or-matrix"
    if len(text) < 90 and shape_count < 12:
        return "section-or-context"
    if any(token in text.lower() for token in [" vs ", "versus", "مقایسه"]):
        return "comparison"
    if any(token in text for token in ["→", "↔", "←", "چرخه", "فرآیند", "گام"]):
        return "process-or-relationship"
    if image_count:
        return "mixed"
    return "structured-text-or-diagram"


def extract_notes(zip_file: zipfile.ZipFile, slide_number: int) -> list[str]:
    rel_name = f"ppt/slides/_rels/slide{slide_number}.xml.rels"
    if rel_name not in zip_file.namelist():
        return []
    from lxml import etree

    rel_root = etree.fromstring(zip_file.read(rel_name))
    notes_target = None
    for rel in rel_root:
        if str(rel.get("Type", "")).endswith("/notesSlide"):
            notes_target = rel.get("Target")
            break
    if not notes_target:
        return []
    base = PurePosixPath("ppt/slides")
    target = str((base / notes_target).as_posix())
    parts = []
    for item in PurePosixPath(target).parts:
        if item == "..":
            if parts:
                parts.pop()
        elif item != ".":
            parts.append(item)
    note_name = "/".join(parts)
    if note_name not in zip_file.namelist():
        return []
    note_root = etree.fromstring(zip_file.read(note_name))
    texts = [normalize_text(v) for v in note_root.xpath(".//a:t/text()", namespaces=XML_NS)]
    return [value for value in texts if value and value not in {str(slide_number)}]


def extract_presentation(deck_path: Path) -> dict:
    prs = Presentation(deck_path)
    slide_width = prs.slide_width
    slide_height = prs.slide_height
    slides = []

    with zipfile.ZipFile(deck_path) as archive:
        members = archive.namelist()
        package_metrics = {
            "slideCount": len([n for n in members if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)]),
            "notesSlideCount": len([n for n in members if re.fullmatch(r"ppt/notesSlides/notesSlide\d+\.xml", n)]),
            "chartCount": len([n for n in members if re.fullmatch(r"ppt/charts/chart\d+\.xml", n)]),
            "embeddingCount": len([n for n in members if n.startswith("ppt/embeddings/") and not n.endswith("/")]),
            "mediaCount": len([n for n in members if n.startswith("ppt/media/") and not n.endswith("/")]),
            "diagramPartCount": len([n for n in members if n.startswith("ppt/diagrams/") and n.endswith(".xml")]),
            "commentPartCount": len([n for n in members if "comment" in n.lower() and n.endswith(".xml")]),
        }

        for number, slide in enumerate(prs.slides, start=1):
            shapes = []
            for shape in slide.shapes:
                shapes.extend(shape_payload(shape, slide_width, slide_height))
            text_items = [(shape.get("text") or {}).get("text") or "" for shape in shapes]
            text_items.extend(
                cell
                for shape in shapes
                for row in (shape.get("table") or {}).get("rows", [])
                for cell in row
            )
            text = normalize_text("\n".join(item for item in text_items if item))
            title = normalize_text(slide.shapes.title.text) if slide.shapes.title and slide.shapes.title.has_text_frame else ""
            if not title:
                title = infer_title(shapes)
            payload = {
                "slideNumber": number,
                "title": title or f"Slide {number}",
                "text": text,
                "textHash": hashlib.sha256(text.encode("utf-8")).hexdigest() if text else None,
                "notes": extract_notes(archive, number),
                "layoutName": slide.slide_layout.name if slide.slide_layout else None,
                "metrics": {
                    "shapeCount": len(shapes),
                    "textShapeCount": sum(1 for shape in shapes if (shape.get("text") or {}).get("text")),
                    "tableCount": sum(1 for shape in shapes if shape.get("table")),
                    "imageCount": sum(1 for shape in shapes if shape.get("image")),
                    "connectorCount": sum(1 for shape in shapes if shape.get("isConnector")),
                    "characterCount": len(text),
                },
                "shapes": shapes,
            }
            payload["contentType"] = classify_slide(payload)
            slides.append(payload)

    duplicate_groups = defaultdict(list)
    for slide in slides:
        if slide["textHash"]:
            duplicate_groups[slide["textHash"]].append(slide["slideNumber"])
    duplicates = [numbers for numbers in duplicate_groups.values() if len(numbers) > 1]
    for slide in slides:
        slide["duplicateGroup"] = next((group for group in duplicates if slide["slideNumber"] in group), None)

    return {
        "source": {
            "filename": deck_path.name,
            "sizeBytes": deck_path.stat().st_size,
            "modifiedUtc": datetime.fromtimestamp(deck_path.stat().st_mtime, tz=timezone.utc).isoformat(),
            "sha256": sha256(deck_path),
        },
        "dimensionsEmu": {"width": slide_width, "height": slide_height},
        "packageMetrics": package_metrics,
        "contentTypeCounts": dict(Counter(slide["contentType"] for slide in slides)),
        "duplicateGroups": duplicates,
        "slides": slides,
    }


def extract_pdf(pdf_path: Path) -> dict:
    pages = []
    with pdfplumber.open(pdf_path) as pdf:
        metadata = dict(pdf.metadata or {})
        for index, page in enumerate(pdf.pages, start=1):
            text = normalize_text(page.extract_text(layout=True) or "")
            pages.append(
                {
                    "pageNumber": index,
                    "text": text,
                    "textHash": hashlib.sha256(text.encode("utf-8")).hexdigest() if text else None,
                    "width": float(page.width),
                    "height": float(page.height),
                }
            )
    return {
        "source": {
            "filename": pdf_path.name,
            "sizeBytes": pdf_path.stat().st_size,
            "modifiedUtc": datetime.fromtimestamp(pdf_path.stat().st_mtime, tz=timezone.utc).isoformat(),
            "sha256": sha256(pdf_path),
        },
        "metadata": metadata,
        "pageCount": len(pages),
        "pages": pages,
    }


def compare_sources(deck: dict, pdf: dict) -> dict:
    def normalize_pdf_token(token: str) -> str:
        if re.search(r"[\u0600-\u06ff]", token):
            return token[::-1]
        return token

    comparisons = []
    for slide, page in zip(deck["slides"], pdf["pages"]):
        ppt_tokens = set(re.findall(r"[\w\u0600-\u06ff]+", slide["text"].casefold()))
        pdf_tokens = {
            normalize_pdf_token(token)
            for token in re.findall(r"[\w\u0600-\u06ff]+", page["text"].casefold())
        }
        union = ppt_tokens | pdf_tokens
        similarity = len(ppt_tokens & pdf_tokens) / len(union) if union else 1.0
        comparisons.append(
            {
                "slideNumber": slide["slideNumber"],
                "pptxCharacters": len(slide["text"]),
                "pdfCharacters": len(page["text"]),
                "tokenJaccard": round(similarity, 4),
            }
        )
    return {
        "samePageCount": len(deck["slides"]) == len(pdf["pages"]),
        "comparedCount": len(comparisons),
        "lowSimilaritySlides": [item for item in comparisons if item["tokenJaccard"] < 0.5],
        "pages": comparisons,
    }


def render_contact_sheets(pdf_path: Path, output_dir: Path, prefix: str, scale: float = 0.32) -> list[str]:
    import pypdfium2 as pdfium

    output_dir.mkdir(parents=True, exist_ok=True)
    document = pdfium.PdfDocument(str(pdf_path))
    thumb_width = 480
    thumb_height = 270
    label_height = 28
    margin = 14
    columns = 3
    rows = 4
    page_width = margin + columns * (thumb_width + margin)
    page_height = margin + rows * (thumb_height + label_height + margin)
    output_paths = []
    font = ImageFont.load_default()

    for sheet_index in range((len(document) + columns * rows - 1) // (columns * rows)):
        sheet = Image.new("RGB", (page_width, page_height), "#20242a")
        draw = ImageDraw.Draw(sheet)
        start = sheet_index * columns * rows
        end = min(start + columns * rows, len(document))
        for page_index in range(start, end):
            slot = page_index - start
            row, column = divmod(slot, columns)
            x = margin + column * (thumb_width + margin)
            y = margin + row * (thumb_height + label_height + margin)
            page = document[page_index]
            bitmap = page.render(scale=scale)
            image = bitmap.to_pil().convert("RGB")
            image.thumbnail((thumb_width, thumb_height), Image.Resampling.LANCZOS)
            frame = Image.new("RGB", (thumb_width, thumb_height), "white")
            offset = ((thumb_width - image.width) // 2, (thumb_height - image.height) // 2)
            frame.paste(image, offset)
            sheet.paste(frame, (x, y))
            draw.text((x, y + thumb_height + 6), f"{prefix} {page_index + 1}", fill="white", font=font)
        output_path = output_dir / f"{prefix.lower()}-contact-{sheet_index + 1:02d}.jpg"
        sheet.save(output_path, quality=88, optimize=True)
        output_paths.append(str(output_path))
    return output_paths


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--deck", type=Path, required=True)
    parser.add_argument("--pdf", type=Path, required=True)
    parser.add_argument("--brandbook", type=Path)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()

    args.out.mkdir(parents=True, exist_ok=True)
    deck = extract_presentation(args.deck)
    pdf = extract_pdf(args.pdf)
    comparison = compare_sources(deck, pdf)

    (args.out / "presentation.json").write_text(json.dumps(deck, ensure_ascii=False, indent=2), encoding="utf-8")
    (args.out / "pdf.json").write_text(json.dumps(pdf, ensure_ascii=False, indent=2), encoding="utf-8")
    (args.out / "source-comparison.json").write_text(json.dumps(comparison, ensure_ascii=False, indent=2), encoding="utf-8")
    rendered = render_contact_sheets(args.pdf, args.out / "contact-sheets", "Slide")
    if args.brandbook:
        rendered.extend(render_contact_sheets(args.brandbook, args.out / "contact-sheets", "Brand"))

    summary = {
        "generatedUtc": datetime.now(timezone.utc).isoformat(),
        "deck": deck["source"],
        "pdf": pdf["source"],
        "packageMetrics": deck["packageMetrics"],
        "contentTypeCounts": deck["contentTypeCounts"],
        "duplicateGroups": deck["duplicateGroups"],
        "comparison": {
            "samePageCount": comparison["samePageCount"],
            "lowSimilarityCount": len(comparison["lowSimilaritySlides"]),
        },
        "contactSheets": rendered,
    }
    (args.out / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
